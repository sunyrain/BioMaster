"""Independent exported-design checks against original source tables."""
import json
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
import design_spr64_512_20260909 as design
import build_retargetmap_spr512_ours_frozen_v2 as source

out=design.OUT
p=pd.read_csv(out/'INTERNAL_MASTER_512.csv')
s=p[p.selection_role.ne('POSITIVE_CONTROL')].copy()
checks={}
identity=json.loads((out/'PREPARATION_IDENTITY.json').read_text())
checks['input_hashes_unchanged']=all(source.sha256(design.ROOT/path)==digest for path,digest in identity['sources'].items())
keys=set(zip(s.ligand_inchikey,s.target_chembl_id))
for name,path,key in [('training',source.TRAINING_PATH,'parent_standard_inchi_key'),('chembl_all_labels',source.STRICT_PATH,'parent_standard_inchi_key'),('kirhub',source.KIRHUB_PATH,'ligand_inchikey')]:
 table=pd.read_csv(path,usecols=[key,'target_chembl_id'])
 checks[name+'_intersection_zero']=not keys.intersection(zip(table[key],table.target_chembl_id))
r=pd.read_csv(source.RANK_PATH)
r['recomputed_rank']=r.groupby('ligand_inchikey',sort=False).independent_validation_rank_score.rank(method='first',ascending=False)
ranks=r.set_index(['ligand_inchikey','target_chembl_id']).recomputed_rank
routed=pd.read_csv(source.SCORE_PATH).set_index(['ligand_inchikey','target_chembl_id']).routed_rank_within_drug
core=s.experiment_arm.eq('FROZEN_384_CORE_DISCOVERY')
actual=[]
for row,is_core in zip(s.itertuples(),core):
 rank=float((ranks if is_core else routed).loc[(row.ligand_inchikey,row.target_chembl_id)])
 actual.append(rank)
s['source_rank']=actual
checks['exported_ranks_match_sources']=all(abs(row.source_rank-(row.retargetmap_rank_384 if is_core else row.routed_rank_within_drug))<1e-8 for row,is_core in zip(s.itertuples(),core))
checks['source_rank_bands_valid']=bool(((s.selection_role.eq('OUR_FROZEN_MODEL_HIGH') & s.source_rank.le(core.map({True:20,False:30}))) | (s.selection_role.eq('OUR_FROZEN_MODEL_INTERMEDIATE') & s.source_rank.between(80,180)) | (s.selection_role.eq('OUR_FROZEN_MODEL_LOW_BACKGROUND') & s.source_rank.gt(250))).all())
dm=s.drop_duplicates('ligand_inchikey')
checks['scaffold_max5_drugs']=bool(dm.groupby('murcko_scaffold').size().le(5).all())
checks['high_scaffold_max2_per_target']=bool(s[s.selection_role.eq('OUR_FROZEN_MODEL_HIGH')].groupby(['gene_symbol','murcko_scaffold']).size().le(2).all())
checks['kinase_drugs_max30']=int(dm.known_kinase_drug.sum())<=30
wb=load_workbook(out/'BLINDED_DESIGN_REVIEW.xlsx',read_only=True)
columns=[c.value for c in next(wb.active.iter_rows())]
checks['blinded_single_sheet_no_internal_columns']=len(wb.sheetnames)==1 and not any(any(t in c for t in ['rank','selection_role','tanimoto','score']) for c in columns)
checks['blinded_512_rows']=wb.active.max_row==513
wb.close()
(out/'INDEPENDENT_VALIDATION.json').write_text(json.dumps(checks,indent=2)+'\n')
print(json.dumps(checks,indent=2))
assert all(checks.values()),'Independent validation failed'
h=s[s.selection_role.eq('OUR_FROZEN_MODEL_HIGH')]
audit=h.groupby('gene_symbol').agg(high_pairs=('pair_id','size'),positive_similarity_median=('positive_max_tanimoto','median'),low_positive_support=('pair_review_status',lambda x:x.eq('LOW_POSITIVE_CHEMICAL_SUPPORT_REVIEW').sum()),close_negative_review=('pair_review_status',lambda x:x.eq('CLOSE_NEGATIVE_REFERENCE_REVIEW').sum()))
audit.to_csv(out/'TARGET_HIGH_SUPPORT_REVIEW.csv')
